from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from typing import BinaryIO

from openpyxl import load_workbook

from backend.app.models.activity_hours import ActivityParticipantHour


@dataclass(frozen=True)
class ActivityHoursExcelConfig:
    """กำหนด worksheet และหัวคอลัมน์สำหรับ IF-XLS-01 และ FR-HRS-02."""

    worksheet_name: str
    header_row: int
    student_reference_header: str
    participant_name_header: str
    reported_hour_type_header: str
    reported_hours_header: str


class ActivityHoursExcelError(ValueError):
    """ข้อผิดพลาดจาก workbook ที่ไม่ตรง config ของ IF-XLS-01 และ FR-HRS-02."""


def read_activity_hours_excel(
    file: BinaryIO,
    config: ActivityHoursExcelConfig,
    document_id: int | None = None,
) -> list[ActivityParticipantHour]:
    """แปลงทุกแถวผู้เข้าร่วมจาก Excel เป็น ActivityParticipantHour ตาม IF-XLS-01 และ FR-HRS-02."""
    if config.header_row < 1:
        raise ValueError("header_row ต้องเป็นเลขแถวเริ่มจาก 1")

    workbook = load_workbook(file, read_only=True, data_only=True)
    try:
        if config.worksheet_name not in workbook.sheetnames:
            raise ActivityHoursExcelError(
                f"ไม่พบ worksheet '{config.worksheet_name}' ในไฟล์ Excel"
            )

        worksheet = workbook[config.worksheet_name]
        header_values = next(
            worksheet.iter_rows(
                min_row=config.header_row,
                max_row=config.header_row,
                values_only=True,
            ),
            None,
        )
        if header_values is None:
            raise ActivityHoursExcelError(
                f"ไม่พบแถวหัวตารางที่ {config.header_row}"
            )

        headers = {
            value: index
            for index, value in enumerate(header_values)
            if value is not None
        }
        configured_headers = {
            "student_reference": config.student_reference_header,
            "participant_name": config.participant_name_header,
            "reported_hour_type": config.reported_hour_type_header,
            "reported_hours": config.reported_hours_header,
        }
        if len(set(configured_headers.values())) != len(configured_headers):
            raise ActivityHoursExcelError("ชื่อหัวคอลัมน์ที่กำหนดต้องไม่ซ้ำกัน")
        missing_headers = [
            header
            for header in configured_headers.values()
            if header not in headers
        ]
        if missing_headers:
            raise ActivityHoursExcelError(
                f"ไม่พบหัวคอลัมน์ที่กำหนด: {', '.join(missing_headers)}"
            )

        records: list[ActivityParticipantHour] = []
        for row_number, values in enumerate(
            worksheet.iter_rows(
                min_row=config.header_row + 1,
                values_only=True,
            ),
            start=config.header_row + 1,
        ):
            selected_values = {
                field: values[headers[header]] if headers[header] < len(values) else None
                for field, header in configured_headers.items()
            }
            if all(value is None or value == "" for value in selected_values.values()):
                continue

            try:
                student_reference = _required_text(
                    selected_values["student_reference"], "student_reference"
                )
                participant_name = _required_text(
                    selected_values["participant_name"], "participant_name"
                )
                reported_hour_type = _required_text(
                    selected_values["reported_hour_type"], "reported_hour_type"
                )
                reported_hours = _required_decimal(selected_values["reported_hours"])
            except ValueError as error:
                raise ActivityHoursExcelError(
                    f"ข้อมูลแถวที่ {row_number} ไม่ถูกต้อง: {error}"
                ) from error

            records.append(
                ActivityParticipantHour(
                    document_id=document_id,
                    student_reference=student_reference,
                    participant_name=participant_name,
                    reported_hour_type=reported_hour_type,
                    reported_hours=reported_hours,
                )
            )

        return records
    finally:
        workbook.close()


def _required_text(value: object, field_name: str) -> str:
    if value is None or not str(value).strip():
        raise ValueError(f"{field_name} ต้องมีค่า")
    return str(value)


def _required_decimal(value: object) -> Decimal:
    if value is None or isinstance(value, bool):
        raise ValueError("reported_hours ต้องเป็นตัวเลข")
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError) as error:
        raise ValueError("reported_hours ต้องเป็นตัวเลข") from error
    if not result.is_finite():
        raise ValueError("reported_hours ต้องเป็นตัวเลขจำกัด")
    return result
